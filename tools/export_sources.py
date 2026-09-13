#!/usr/bin/env python3
"""Export the historical 41-source catalogue and its correspondence to OpenPlant.

Inputs are read-only. This tool requires openpyxl and writes only the four source
documentation outputs below -- never the input spreadsheets or original code.
DOI checks are optional and use Crossref's public API (no API key required).
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import html
import json
import re
import urllib.parse
import urllib.request
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

# Order is the order of the 41 data rows in the original Sheet4 catalogue.
# source_id is exactly the corresponding worksheet name in the release mapping.
SOURCE_KEYS = [
    ("PlantVillage", "hughesOpenAccessRepository2016"),
    ("OLID I", "orkaOLIDOpenLeaf2023"),
    ("PlantDoc", "singhPlantDocDatasetVisual2020"),
    ("RiceDID", "huyminhdoRice2019"),
    ("CucumberPDD", "karimnegmCucumber2020"),
    ("WheatLD", "getachewWheat2021"),
    ("CGIAR", "shadabhussainCGIAR2020"),
    ("AppleLeaf9", "yangEfficientIdentificationApple2022"),
    ("BananaLDI", "hailuBanana2021"),
    ("ATLDSD", "fengApple2022"),
    ("chilli", "naikDetection2022"),
    ("CottonLDD", "noonComputationally2021"),
    ("CottonPD", "dhamodharanrCotton2023"),
    ("PLD", "rashidMultilevelDeepLearning2021a"),
    ("Cassava", "mwebazeICassava2019Finegrained2019"),
    ("PP2020", "thapaPlant2020"),
    ("ESCA", "alessandriniGrapevine2021"),
    ("SugarcaneLDD", "daphalSugarcane2022"),
    ("CornLD", "sandiindikasaputraCorn2023"),
    ("MangoLeafBD", "ahmedMangoLeafBD2023"),
    ("Soybean Leaves", "mignoniSoybean2022"),
    ("SugarcaneDD", "prabhakaransoundarSugarcane2022"),
    ("Coffee plant disease", "coffeediseaseCoffee2021"),
    ("Fruit Leaf", "chouhanDataRepositoryLeaf2019"),
    ("Weed25", "wangWeed25DeepLearning2022"),
    ("CottonWeedDet12", "dangYOLOWeeds2023"),
    ("DeepWeeds", "olsenDeepWeeds2019"),
    ("CWID15", "chenPerformance2022"),
    ("WeedNet-R", "guoWeedNetRSugarBeet2023"),
    ("CWD30", "ilyasCWD302025"),
    ("PlantSeedlings", "giselssonPublicImageDatabase2017"),
    ("plantnet_300K", "garcinPlntNet300KImageDataset2021"),
    ("Early_crop_weed", "auagroupAUAgroup2024"),
    ("OPPD", "leminenmadsenOpen2020"),
    ("CottonWeedDet3", "rahmanPerformance2023"),
    ("weed-datasets", "zhangchuanyinZhangchuanyin2023"),
    ("Carrot-Weed", "lameskiWeedDetectionDataset2017"),
    ("VCD", "lacAnnotated2022"),
    ("SorghumWeed", "justinaSorghumWeedDataset_Classification2024"),
    ("ImageWeeds", "raiMultiformat2023"),
    ("TobSet", "alamTobSet2022"),
]

# Facts below were checked on the linked primary dataset pages on this date.
REVIEW_DATE = "2026-09-12"
LICENCES = {
    "PlantVillage": "CC0-1.0",  # Posted for the recorded Mendeley deposit.
    "WheatLD": "CC-BY-4.0", "BananaLDI": "CC-BY-4.0",
    "chilli": "CC-BY-4.0", "ESCA": "CC-BY-4.0",
    "SugarcaneLDD": "CC-BY-4.0", "MangoLeafBD": "CC-BY-NC-3.0",
    "Soybean Leaves": "CC-BY-4.0", "Fruit Leaf": "CC-BY-4.0",
    "VCD": "CC-BY-4.0", "SorghumWeed": "CC-BY-4.0",
    "PlantSeedlings": "CC-BY-SA-4.0", "OPPD": "CC-BY-NC-SA-4.0",
    "ImageWeeds": "CC-BY-4.0", "plantnet_300K": "CC-BY-4.0",
}
LICENCE_EVIDENCE = {
    "ImageWeeds": "https://api.datacite.org/dois/10.17632/8kjcztbjz2.2",
    "plantnet_300K": "https://api.datacite.org/dois/10.5281/zenodo.5645731",
}
DATA_DOIS = {
    "PlantVillage": "10.17632/tywbtsjrjv.1", "WheatLD": "10.17632/wgd66f8n6h.1",
    "BananaLDI": "10.17632/rjykr62kdh.1", "chilli": "10.17632/tf9dtfz9m6.2",
    "ESCA": "10.17632/89cnxc58kj.1", "SugarcaneLDD": "10.17632/9424skmnrk.1",
    "MangoLeafBD": "10.17632/hxsnvwty3r.1", "Soybean Leaves": "10.17632/bycbh73438.1",
    "Fruit Leaf": "10.17632/hb74ynkjcn.4", "VCD": "10.17632/d7kbzjr83k.1",
    "SorghumWeed": "10.17632/4gkcyxjyss.1", "plantnet_300K": "10.5281/zenodo.5645731",
    "ImageWeeds": "10.17632/8kjcztbjz2.2",
}
NOTES = {
    "PlantVillage": "The original adapter uses Plant_leave_diseases_dataset_with_augmentation. The recorded Mendeley deposit reports 61,486 augmented images; the catalogue's 54,309 describes PlantVillage generally. Reconstruct from the recorded deposit and manifest, not an interchangeable PlantVillage mirror. The displayed CC0 label is the deposit's declaration, not an independent clearance of upstream image rights.",
    "CGIAR": "The spreadsheet's paper URL mistakenly points to the AppleLeaf9 paper. This catalogue uses the CGIAR data-page citation from Table 1 and the manuscript bibliography.",
    "chilli": "The recorded download is Mendeley V2 (2024), deposited by Aishwarya M P and Padmanabha Reddy. The manuscript cites the Naik et al. (2022) article. The download-to-paper relationship is not established by the sparse deposit description; both are retained as recorded provenance.",
    "PLD": "The historical spreadsheet reports 4,072 upstream images; the manuscript Table 1 reports 4,062. The released mapping independently records 1,020 selected images.",
    "Cassava": "The retained URL and cited paper are the 2019 iCassava challenge. The catalogue reports 21,400 upstream images. An exact upstream archive version was not recorded; the manifest's filenames and selected 316 images define this release's membership.",
    "WeedNet-R": "SugarBeet2016 in the catalogue corresponds to the WeedNet-R adapter and worksheet. The WeedNet-R repository acknowledges Sugar Beets 2016 as its upstream source. OpenPlant samples are object crops and can outnumber original images.",
    "PlantSeedlings": "The adapter reads NonsegmentedV2 (cropped, nonsegmented single plants). Use the V2 link on the official dataset page.",
    "plantnet_300K": "The manuscript cites Zenodo record 5645731 (V1.1), whose DataCite record declares CC BY 4.0 at deposit level. The official repository also supplies the NeurIPS 2021 paper and per-image author/license metadata. Preserve those image-level terms. The 1,081 upstream labels correspond to 1,021 distinct class names in the released mapping after name standardization.",
    "OPPD": "The upstream site describes 7,590 RGB images containing 315,038 plant objects. The catalogue's 315,038 therefore counts objects. OpenPlant uses selected individual-plant images recorded under images_plants.",
    "CottonWeedDet3": "Listed in the historical catalogue and manuscript Table 1, but no separately named worksheet or adapter registration occurs in the recovered release mapping/build registry. No contribution can be assigned from that evidence; the reason is unresolved.",
    "Carrot-Weed": "Listed in the historical catalogue and manuscript Table 1, but no separately named worksheet or adapter registration occurs in the recovered release mapping/build registry. No contribution can be assigned from that evidence; the reason is unresolved.",
    "SorghumWeed": "Only the 1,404 sorghum samples are represented in the mapping; the upstream generic grasses/broad-leaf weed classes are not represented as separate OpenPlant taxa.",
    "ImageWeeds": "The recorded Mendeley V2 link returned a page without dataset metadata during review; its DataCite DOI record confirms the deposit and CC BY 4.0 declaration. Use the linked Data in Brief article for access information if the deposit page is unavailable.",
    "Weed25": "The paper supplies the Baidu download link and access code rn5h. The historical catalogue used the paper DOI as its data URL; both routes are retained here.",
}

EXTRA_BIB = r"""
@inproceedings{garcinPlantNetPaper2021,
  title = {{Pl@ntNet-300K}: a plant image dataset with high label ambiguity and a long-tailed distribution},
  author = {Garcin, Camille and Joly, Alexis and Bonnet, Pierre and Lombardo, Jean-Christophe and Affouard, Antoine and Chouet, Mathias and Servajean, Maximilien and Lorieul, Titouan and Salmon, Joseph},
  year = {2021},
  booktitle = {NeurIPS Datasets and Benchmarks},
  url = {https://datasets-benchmarks-proceedings.neurips.cc/paper/2021/file/7e7757b1e12abcb736ab9a754ffb617a-Paper-round2.pdf}
}
@misc{pandianPlantVillageDeposit2019,
  title = {Data for: Identification of Plant Leaf Diseases Using a 9-layer Deep Convolutional Neural Network},
  author = {{ARUN PANDIAN J} and {GEETHARAMANI GOPAL}},
  year = {2019},
  publisher = {Mendeley Data},
  doi = {10.17632/tywbtsjrjv.1},
  url = {https://data.mendeley.com/datasets/tywbtsjrjv/1},
  version = {1}
}
@misc{aishwaryaChilliDeposit2024,
  title = {chilli dataset},
  author = {{Aishwarya M P} and {Padmanabha Reddy}},
  year = {2024},
  publisher = {Mendeley Data},
  doi = {10.17632/tf9dtfz9m6.2},
  url = {https://data.mendeley.com/datasets/tf9dtfz9m6/2},
  version = {2}
}
"""


def parse_bib(text):
    """Read the one-field-per-line BibTeX form used by the source bibliography."""
    entries = {}
    for block in re.split(r"(?m)(?=^@)", text):
        match = re.match(r"@(\w+)\{([^,]+),", block)
        if not match:
            continue
        fields = dict(re.findall(r"(?m)^\s+(\w+)\s*=\s*\{(.*)\},?\s*$", block))
        fields.update(citekey=match[2], bibtype=match[1], raw=block.strip())
        entries[match[2]] = fields
    return entries


def plain(value):
    value = html.unescape(value)
    for old, new in [(r"\_", "_"), (r"\o", "ø"), (r"\relax ", ""),
                     (r"\emph", ""), (r'\"u', 'ü'), (r'\c s', 'ş'),
                     (r"\'e", "é")]:
        value = value.replace(old, new)
    return re.sub(r"\s+", " ", value.replace("{", "").replace("}", "")).strip()


def normalized(value):
    return re.sub(r"[^a-z0-9]", "", plain(value).lower())


def crossref_check(entry):
    doi = entry.get("doi", "")
    if not doi or doi.startswith(("10.48550/", "10.5281/", "10.11922/", "10.17632/")):
        return None
    url = "https://api.crossref.org/works/" + urllib.parse.quote(doi, safe="/")
    try:
        request = urllib.request.Request(url, headers={"User-Agent": "OpenPlant-source-catalogue/1.0"})
        with urllib.request.urlopen(request, timeout=25) as response:
            result = json.load(response)["message"]
        title = result.get("title", [""])[0]
        if result.get("subtitle"):
            title += ": " + result["subtitle"][0]
        similarity = SequenceMatcher(None, normalized(title), normalized(entry["title"])).ratio()
        years = {date[0] for field in ("published", "published-print", "published-online", "issued")
                 for date in result.get(field, {}).get("date-parts", []) if date}
        ok = similarity > 0.9 and int(entry["year"]) in years
        return {"status": "verified" if ok else "mismatch", "url": url,
                "retrieved_title": title, "retrieved_years": sorted(years),
                "title_similarity": round(similarity, 4), "checked_on": REVIEW_DATE}
    except Exception as error:
        return {"status": "manual_needed", "url": url,
                "note": f"Crossref lookup failed: {type(error).__name__}", "checked_on": REVIEW_DATE}


def reference(entry, checks):
    doi = entry.get("doi", "")
    url = entry.get("url") or ("https://doi.org/" + doi if doi else entry.get("howpublished", ""))
    if doi.startswith("10.48550/arXiv."):
        status = {"status": "verified", "url": "https://arxiv.org/abs/" + doi.split("arXiv.")[1],
                  "checked_on": REVIEW_DATE, "note": "Title, authors and version year checked on arXiv."}
    elif entry["citekey"] in {"garcinPlntNet300KImageDataset2021", "garcinPlantNetPaper2021",
                              "pandianPlantVillageDeposit2019", "aishwaryaChilliDeposit2024",
                              "getachewWheat2021", "hailuBanana2021", "daphalSugarcane2022"}:
        status = {"status": "verified", "url": url, "checked_on": REVIEW_DATE,
                  "note": "Bibliographic metadata checked on the primary repository or dataset page."}
    else:
        status = checks.get(entry["citekey"]) or {
            "status": "manuscript_bibliography", "url": url,
            "note": "Transcribed from the manuscript bibliography; not independently verified here."}
    return {"citekey": entry["citekey"], "type": entry["bibtype"],
            "title": plain(entry["title"]), "authors": plain(entry.get("author", "")),
            "year": int(entry["year"]), "venue": plain(entry.get("journal") or entry.get("booktitle") or entry.get("publisher", "")),
            "doi": doi, "url": url, "verification": status}


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_counts(path):
    result = {}
    all_classes = set()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    for sheet in workbook:
        if sheet.title == "datasets":
            continue
        count = Counter()
        classes = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            classes.add(row[0])
            count[str(row[3])] += 1
        assert set(count) <= {"train", "val", "test"}, (sheet.title, count)
        result[sheet.title] = {"samples": sum(count.values()), "classes": len(classes),
                               **{split: count[split] for split in ("train", "val", "test")}}
        all_classes.update(classes)
    workbook.close()
    assert sum(row["samples"] for row in result.values()) == 635176
    assert len(all_classes) == 1167
    assert len(result) == 39
    return result


def export(args):
    entries = parse_bib(args.bibliography.read_text(encoding="utf-8") + "\n" + EXTRA_BIB)
    selected_keys = [key for _, key in SOURCE_KEYS] + list(parse_bib(EXTRA_BIB))
    assert len(set(selected_keys)) == len(selected_keys)
    selected = {key: entries[key] for key in selected_keys}
    checks = {}
    if args.verify_dois:
        with ThreadPoolExecutor(max_workers=4) as pool:
            checks = {key: result for key, result in zip(selected, pool.map(crossref_check, selected.values())) if result}
        print("DOI checks:", dict(Counter(check["status"] for check in checks.values())), flush=True)
        for key, check in checks.items():
            if check["status"] != "verified":
                print(key, json.dumps(check, ensure_ascii=True), flush=True)
    counts = read_counts(args.mapping)
    catalogue = openpyxl.load_workbook(args.catalogue, read_only=True, data_only=True)
    rows = [row for row in catalogue["Sheet4"].iter_rows(min_row=2, values_only=True) if row[1]]
    assert len(rows) == len(SOURCE_KEYS) == 41
    records = []
    for index, (row, (source_id, citekey)) in enumerate(zip(rows, SOURCE_KEYS), start=1):
        recorded_data_url = str(row[12]).strip()
        data_url = {
            "plantnet_300K": "https://zenodo.org/records/5645731",
            "Weed25": "https://pan.baidu.com/s/1rnUoDm7IxxmX1n1LmtXNXw",
            "CWD30": "https://github.com/Mr-TalhaIlyas/CWD30",
        }.get(source_id, recorded_data_url)
        extra_keys = {"PlantVillage": ["pandianPlantVillageDeposit2019"],
                      "chilli": ["aishwaryaChilliDeposit2024"],
                      "plantnet_300K": ["garcinPlantNetPaper2021"]}.get(source_id, [])
        details = counts.get(source_id)
        record = {
            "catalogue_index": index, "source_id": source_id,
            "slug": re.sub(r"[^a-z0-9]+", "-", source_id.lower()).strip("-"),
            "name": str(row[1]),
            "source_group": "disease" if index <= 24 else "general_plant" if source_id == "plantnet_300K" else "crop_weed",
            "catalogue_task": row[2], "catalogue_class_count": row[3],
            "catalogue_species_count": row[4], "catalogue_image_count": row[5],
            "observation_scale": row[6], "environment": row[7], "target_species": row[8],
            "data_url": data_url, "data_doi": DATA_DOIS.get(source_id, ""),
            "homepage_url": recorded_data_url, "recorded_data_url": recorded_data_url,
            "mapping_sheet": source_id if details else "",
            "release_status": "mapped" if details else "catalogue_only",
            "mapped_samples": details["samples"] if details else None,
            "mapped_classes": details["classes"] if details else None,
            "train": details["train"] if details else None,
            "val": details["val"] if details else None,
            "test": details["test"] if details else None,
            "dataset_license": LICENCES.get(source_id, "unknown"),
            "license_status": "declared_in_datacite_record" if source_id in LICENCE_EVIDENCE else "declared_on_dataset_page" if source_id in LICENCES else "unknown",
            "license_evidence_url": LICENCE_EVIDENCE.get(source_id, data_url if source_id in LICENCES else ""),
            "license_checked_on": REVIEW_DATE if source_id in LICENCES else "",
            "references": [reference(entries[key], checks) for key in [citekey] + extra_keys],
            "notes": NOTES.get(source_id, ""),
        }
        records.append(record)
    catalogue.close()
    assert set(counts) == {r["source_id"] for r in records if r["release_status"] == "mapped"}
    out = args.output
    for name in ("metadata", "references", "docs"):
        (out / name).mkdir(parents=True, exist_ok=True)
    payload = {"schema_version": 1, "review_date": REVIEW_DATE,
               "catalogue_sources": 41, "mapped_sources": 39,
               "mapped_samples": 635176, "mapped_classes": 1167,
               "input_provenance": {name: {"filename": path.name, "sha256": file_hash(path)}
                                    for name, path in [("catalogue", args.catalogue), ("mapping", args.mapping), ("bibliography", args.bibliography)]},
               "sources": records}
    (out / "metadata/sources.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    fields = [key for key in records[0] if key != "references"] + [
        "citekey", "reference_type", "reference_title", "reference_authors", "reference_year",
        "reference_venue", "reference_doi", "reference_url", "reference_verification", "additional_citekeys"]
    with (out / "metadata/sources.csv").open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        for record in records:
            flat = {key: value for key, value in record.items() if key != "references"}
            ref = record["references"][0]
            flat.update(citekey=ref["citekey"], reference_type=ref["type"], reference_title=ref["title"],
                        reference_authors=ref["authors"], reference_year=ref["year"], reference_venue=ref["venue"],
                        reference_doi=ref["doi"], reference_url=ref["url"],
                        reference_verification=ref["verification"]["status"],
                        additional_citekeys=";".join(item["citekey"] for item in record["references"][1:]))
            writer.writerow(flat)
    bib = "% Source references for OpenPlant. Article, preprint, repository and dataset entries retain distinct types.\n\n"
    for key in sorted(selected):
        entry = selected[key]
        raw = entry["raw"]
        # Retain every original field while adding a portable URL where absent.
        if "url" not in entry:
            url = reference(entry, checks)["url"]
            if url:
                raw = raw.rstrip()[:-1].rstrip().rstrip(",") + f",\n  url = {{{url}}}\n}}"
        bib += raw + "\n\n"
    (out / "references/datasets.bib").write_text(bib.rstrip() + "\n", encoding="utf-8")
    write_docs(records, out / "docs/sources.md")
    print(json.dumps({"catalogue_sources": len(records), "mapped_sources": len(counts),
                      "samples": sum(c["samples"] for c in counts.values()),
                      "references": len(selected), "declared_dataset_licenses": len(LICENCES)}, ensure_ascii=True))


def write_docs(records, path):
    text = """# Dataset sources and composition

OpenPlant contains **635,176 samples in 1,167 classes**. Its catalogue lists
**41 datasets**: 24 disease datasets, 16 crop/weed datasets, and one general plant
dataset. The image manifest traces the selected samples to **39 source groups**.

**We cannot currently distribute the complete OpenPlant image collection because
of potential copyright and licensing conflicts.** Some source datasets may prohibit
redistribution of original or modified images. Obtain the data from the providers
below under their respective terms, then follow the
[reconstruction guide](reconstruction.md) to build OpenPlant locally.
See [data licensing](../DATA_LICENSE.md) for the release policy.

[Image manifest](../metadata/manifest.csv.gz) · [Class list](../metadata/classes.csv) ·
[Sources CSV](../metadata/sources.csv) / [JSON](../metadata/sources.json) ·
[Full bibliography](../references/datasets.bib)

## How the collection is formed

Preparation selects healthy examples from disease datasets, screens poorly visible
plants, standardizes species names, and crops annotated plant regions. The manifest
fixes the retained samples and their train/validation/test assignments.

Counts below refer to output classification images. One photograph can yield
several crops, and classes can occur in several sources. Use the manifest totals
for benchmark comparisons and the documented source versions for reconstruction.

## Source catalogue

`source_id` joins this table to the manifest. CottonWeedDet3 and Carrot-Weed have
`catalogue_only` status: both appear in the paper, but neither has a separate
mapping worksheet or registered adapter. Their contributions remain unresolved
and are shown as **—**.

| # | Source dataset | source_id | Mapped samples | Mapped classes | Access | Citation |
|---:|---|---|---:|---:|---|---|
"""
    for r in records:
        ref = r["references"][0]
        samples = f'{r["mapped_samples"]:,}' if r["mapped_samples"] is not None else "—"
        classes = str(r["mapped_classes"]) if r["mapped_classes"] is not None else "—"
        text += f'| {r["catalogue_index"]} | {r["name"]} | `{r["source_id"]}` | {samples} | {classes} | [Data]({r["data_url"]}) | [{ref["year"]}](#{r["slug"]}) |\n'
    text += """
## References, versions and dataset terms

References include papers, preprints, data deposits and repositories. Full author
lists are in the bibliography; verification records are in the JSON.

License labels below reproduce the source metadata reviewed on 12 September 2026.
`unknown` means no data license is recorded there. Check the image terms applicable
to the source version you obtain and retain attribution; a paper or code license
alone does not establish image-reuse rights.
The CWD30 and PlantNet-300K entries include additional terms to check before reuse.

"""
    # Presentation-only notes: these do not change the source metadata or parser.
    notes = {
        "PlantVillage": "Use the recorded augmented deposit (61,486 images), read by the original adapter as `Plant_leave_diseases_dataset_with_augmentation`. The catalogue's 54,309 refers to PlantVillage generally. CC0 is the deposit's declaration; upstream image rights still need checking.",
        "CGIAR": "The spreadsheet linked the AppleLeaf9 paper here. This entry uses the CGIAR data citation from the manuscript.",
        "chilli": "The download is Mendeley V2 (2024), deposited by Aishwarya M P and Padmanabha Reddy. The manuscript cites Naik et al. (2022); the deposit does not establish that connection. Both references are retained.",
        "PLD": "Upstream totals differ: 4,072 in the spreadsheet and 4,062 in the paper. The manifest contains 1,020 selected images.",
        "Cassava": "The link and paper describe the 2019 iCassava challenge; the catalogue lists 21,400 images without an archive version. Use the manifest's filenames for the 316 selected images.",
        "WeedNet-R": "SugarBeet2016 uses the WeedNet-R adapter and worksheet. The repository credits Sugar Beets 2016 as its source; OpenPlant counts the extracted plant crops.",
        "CWD30": "The [official Terms of Use](https://cwd-30.github.io/cwd-30/terms_of_use.html) declare CC BY-NC-SA 4.0 and separately prohibit redistribution of original or modified data. See [data licensing](../DATA_LICENSE.md) for this conflict.",
        "PlantSeedlings": "Use the official page's **NonsegmentedV2** download, as specified by the adapter.",
        "plantnet_300K": "Zenodo V1.1 declares CC BY 4.0 at deposit level. The official repository provides per-image author and license metadata: retain it and check each image's terms. Its 1,081 upstream labels map to 1,021 distinct class names in OpenPlant.",
        "OPPD": "The upstream total of 315,038 counts plant objects in 7,590 RGB photographs. OpenPlant selects individual-plant images from `images_plants`.",
        "CottonWeedDet3": "",
        "Carrot-Weed": "",
        "SorghumWeed": "The manifest includes the 1,404 sorghum images. Generic grass and broad-leaf weed labels are omitted.",
        "ImageWeeds": "DataCite confirms the V2 deposit and CC BY 4.0 declaration. If the Mendeley page is unavailable, consult the linked paper for access information.",
        "Weed25": "Baidu access code: **rn5h**, as supplied in the paper. The recorded access page is the paper DOI.",
    }
    for r in records:
        text += f'<a id="{r["slug"]}"></a>\n\n### {r["name"]}\n\n'
        text += f'[Data]({r["data_url"]}) · License: `{r["dataset_license"]}`'
        if r["data_doi"]:
            text += f' · [Dataset DOI](https://doi.org/{r["data_doi"]})'
        if r["homepage_url"] != r["data_url"]:
            text += f' · [Project / access page]({r["homepage_url"]})'
        if r["license_evidence_url"] and r["license_evidence_url"] != r["data_url"]:
            text += f' · [License declaration]({r["license_evidence_url"]})'
        text += "\n\n"
        for ref in r["references"]:
            authors = ref["authors"].split(" and ")
            author_label = authors[0].split(",")[0] + " et al." if len(authors) > 2 else " and ".join(author.split(",")[0] for author in authors)
            venue = f' {ref["venue"].rstrip(".")}.' if ref["venue"] else ""
            text += f'{author_label} ({ref["year"]}). [{ref["title"]}]({ref["url"]}).{venue} `{ref["citekey"]}`.\n\n'
        note = notes.get(r["source_id"], r["notes"])
        if note:
            text += note + "\n\n"
    text += """## Provenance

Counts come from the 39 source worksheets in `image_mapping_test_0220.xlsx`.
Sheet4 of `datasets(已自动还原).xlsx` supplies the catalogue; `openplant.bib`
supplies the manuscript references, supplemented by primary-source records.
The JSON includes their SHA-256 hashes. `catalogue_image_count` stores the
historical upstream totals; `mapped_samples` counts OpenPlant's selected images.

Regenerate the source files from those inputs:

```bash
python tools/export_sources.py --catalogue /path/to/catalogue.xlsx \\
  --mapping /path/to/image_mapping_test_0220.xlsx \\
  --bibliography /path/to/openplant.bib --output . --verify-dois
```

`--verify-dois` checks article DOIs through Crossref. Data-license declarations
have separate evidence links in the source records.
"""
    path.write_text(text, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--catalogue", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, required=True)
    parser.add_argument("--bibliography", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=Path("."))
    parser.add_argument("--verify-dois", action="store_true")
    export(parser.parse_args())


if __name__ == "__main__":
    main()
